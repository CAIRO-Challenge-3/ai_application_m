from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from PIL import Image

P=Presentation(); P.slide_width=Inches(13.333); P.slide_height=Inches(7.5)
blank=P.slide_layouts[6]
NAVY='10234A'; BLUE='2259D6'; GREEN='15845F'; YELLOW='FFC83D'; RED='D94249'; WHITE='FFFFFF'; INK='14254A'; MUTED='51627D'; PALE='F2F6FF'; LIGHT='E8EFFD'
def rgb(c): return RGBColor.from_string(c)
def shape(s,x,y,w,h,color,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
 a=s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); a.fill.solid(); a.fill.fore_color.rgb=rgb(color); a.line.fill.background() if not line else None
 if line: a.line.color.rgb=rgb(line)
 return a
def txt(s,text,x,y,w,h,size=22,color=INK,bold=False,align=PP_ALIGN.LEFT):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.03); tf.margin_right=Inches(.03); tf.margin_top=Inches(.02); tf.margin_bottom=Inches(.02)
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color); p.alignment=align
 return box
def base(title,sub=None):
 s=P.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(WHITE)
 shape(s,0,0,13.333,.15,BLUE,MSO_SHAPE.RECTANGLE); txt(s,title,.65,.43,12,0.65,31,NAVY,True)
 if sub: txt(s,sub,.68,1.16,12,.5,18,MUTED)
 txt(s,str(len(P.slides)),12.55,7.08,.3,.2,10,MUTED,align=PP_ALIGN.RIGHT)
 return s
def circle(s,x,y,d,color): return shape(s,x,y,d,d,color,MSO_SHAPE.OVAL)
def robot(s,x,y,scale=1):
 shape(s,x+.25*scale,y+.18*scale,1.7*scale,1.3*scale,'DAE8FF'); shape(s,x+.38*scale,y+.32*scale,1.44*scale,.85*scale,WHITE)
 circle(s,x+.62*scale,y+.55*scale,.16*scale,BLUE); circle(s,x+1.42*scale,y+.55*scale,.16*scale,BLUE)
 shape(s,x+.83*scale,y+.89*scale,.55*scale,.06*scale,BLUE,MSO_SHAPE.RECTANGLE)
 shape(s,x+.45*scale,y+1.43*scale,1.3*scale,1.38*scale,BLUE); shape(s,x+.64*scale,y+1.67*scale,.92*scale,.68*scale,WHITE)
 txt(s,'♥',x+.85*scale,y+1.75*scale,.52*scale,.42*scale,int(23*scale),RED,True,PP_ALIGN.CENTER)
 shape(s,x+.06*scale,y+1.55*scale,.35*scale,1.0*scale,'DAE8FF'); shape(s,x+1.8*scale,y+1.55*scale,.35*scale,1.0*scale,'DAE8FF')
 shape(s,x+.69*scale,y+2.76*scale,.85*scale,.24*scale,'DAE8FF')
def card(s,x,y,w,h,heading,body,color=BLUE):
 shape(s,x,y,w,h,PALE); shape(s,x,y,.12,h,color,MSO_SHAPE.RECTANGLE); txt(s,heading,x+.28,y+.25,w-.5,.5,23,color,True); txt(s,body,x+.28,y+.95,w-.55,h-1.05,21,INK)
def logo(s,x,y,w):
 p=s.shapes.add_picture('logo-cropped.png', Inches(x), Inches(y), width=Inches(w)); p.click_action.hyperlink.address='https://www.mainfraenkische-werkstaetten.de/'

s=P.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(NAVY)
shape(s,8.9,-.6,5.3,8.4,BLUE,MSO_SHAPE.OVAL); circle(s,9.25,1.35,3.4,'DCEAFF'); robot(s,9.82,1.65,1.05)
txt(s,'YOUR HELPER\nAT WORK',.8,1.05,7.5,1.8,42,WHITE,True); txt(s,'A safer day.\nA little more independence.',.85,3.15,6.9,1.3,27,'DCEAFF')
shape(s,.85,5.42,4.1,1.15,WHITE); logo(s,1.18,5.48,3.45); txt(s,'Tap the logo to visit the website',.87,6.68,5,.35,13,'DCEAFF')

s=base('Meet your workplace helper','Two simple ways to get support.')
card(s,.75,2.0,5.65,3.65,'1   SAFETY CHECK','The camera checks for\nsafety gear.\n\nNeed something?\nYou get a friendly reminder.',BLUE)
card(s,6.8,2.0,5.65,3.65,'2   WORKSHOP MAP','Tap the map to see more.\n\nTap a symbol to learn\nwhat it means.',GREEN)

s=base('Ready for work?','A quick check can help you remember your safety gear.')
shape(s,.8,2.05,3.3,3.85,PALE); circle(s,1.78,2.4,1.25,'BFD4FF'); shape(s,1.39,3.78,2.05,1.42,BLUE); txt(s,'YOU',1.6,5.34,1.6,.35,16,NAVY,True,PP_ALIGN.CENTER)
txt(s,'→',4.25,3.33,.9,.7,37,BLUE,True,PP_ALIGN.CENTER)
shape(s,5.35,2.05,3.1,3.85,'EAF8F1'); txt(s,'CAMERA',5.65,2.4,2.5,.45,20,GREEN,True,PP_ALIGN.CENTER); shape(s,6.1,3.03,1.6,1.1,GREEN); circle(s,6.56,3.27,.65,WHITE); circle(s,6.72,3.43,.32,GREEN); txt(s,'Looks for gear',5.65,4.65,2.5,.5,20,INK,True,PP_ALIGN.CENTER)
txt(s,'→',8.6,3.33,.9,.7,37,BLUE,True,PP_ALIGN.CENTER)
shape(s,9.65,2.05,2.85,3.85,'FFF6DC'); txt(s,'REMINDER',9.85,2.4,2.45,.45,20,INK,True,PP_ALIGN.CENTER); txt(s,'Please put\non your\nsafety gear.',9.9,3.18,2.35,1.65,24,INK,True,PP_ALIGN.CENTER)
txt(s,'The AI can make mistakes. A person can always help.',.9,6.46,11.8,.42,17,MUTED)

s=base('Find your way','Tap the map. It gets bigger.')
s.shapes.add_picture('layout-1.png', Inches(.75), Inches(1.85), height=Inches(4.82))
shape(s,6.0,2.0,6.4,3.8,PALE); txt(s,'1',6.35,2.3,.55,.6,30,BLUE,True); txt(s,'Tap the workshop map.',7.05,2.33,4.9,.6,24,INK,True)
txt(s,'2',6.35,3.32,.55,.6,30,BLUE,True); txt(s,'Zoom in to see more.',7.05,3.35,4.9,.6,24,INK,True)
txt(s,'3',6.35,4.34,.55,.6,30,BLUE,True); txt(s,'Tap a symbol.',7.05,4.37,4.9,.6,24,INK,True)
txt(s,'Example plan supplied for the challenge.',.8,6.88,7,.28,12,MUTED)

s=base('Tap a symbol. Learn what it means.','See it. Read it. Hear it.')
shape(s,.8,2.0,3.3,3.95,'E9F7F0'); circle(s,1.65,2.47,1.6,GREEN); txt(s,'+',1.86,2.48,1.18,1.07,58,WHITE,True,PP_ALIGN.CENTER); txt(s,'FIRST AID',1.19,4.54,2.5,.5,20,GREEN,True,PP_ALIGN.CENTER)
txt(s,'→',4.28,3.3,.7,.6,34,BLUE,True,PP_ALIGN.CENTER)
shape(s,5.15,2.0,7.35,3.95,PALE); txt(s,'First aid',5.6,2.38,6.2,.55,28,NAVY,True); txt(s,'You can get help here\nif you are hurt.',5.6,3.18,6.2,1.15,25,INK)
shape(s,5.6,4.75,2.65,.65,BLUE); txt(s,'▶  Listen again',5.76,4.86,2.32,.38,18,WHITE,True,PP_ALIGN.CENTER)
txt(s,'Example symbol. Real locations must be checked on site.',.85,6.55,11.8,.37,16,MUTED)

s=base('Colours help you choose','Each colour has a meaning. Tap a colour to explore.')
items=[('BLUE','You must do this.',BLUE,'!'),('RED','Do not do this.',RED,'−'),('GREEN','Help and safe exits.',GREEN,'+'),('YELLOW','Be careful.',YELLOW,'!')]
for i,(name,body,col,icon) in enumerate(items):
 x=.75+i*3.16; shape(s,x,2.1,2.88,3.8,PALE); circle(s,x+.79,2.45,1.3,col); txt(s,icon,x+.91,2.51,1.05,.9,45,NAVY if name=='YELLOW' else WHITE,True,PP_ALIGN.CENTER); txt(s,name,x+.2,4.18,2.48,.45,21,INK,True,PP_ALIGN.CENTER); txt(s,body,x+.22,4.8,2.44,.8,19,INK,False,PP_ALIGN.CENTER)

s=base('A little help. A big difference.','Our idea for a friendlier, safer workplace.')
robot(s,.85,2.05,1.2)
shape(s,4.0,2.05,8.35,3.8,PALE); txt(s,'✓   Remember your safety gear',4.48,2.45,7.35,.6,25,INK,True); txt(s,'✓   Explore the workshop map',4.48,3.4,7.35,.6,25,INK,True); txt(s,'✓   Learn with pictures and sound',4.48,4.35,7.35,.6,25,INK,True)
txt(s,'Thank you!',4.5,6.25,4.5,.7,29,BLUE,True); logo(s,9.35,6.3,3.0)
P.save('Workplace_Helper_Presentation.pptx')

